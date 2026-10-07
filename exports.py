"""Turns resume / tool text into ATS-friendly PDF and Word files (single column, plain fonts, no tables or images)."""
import io, re
from xml.sax.saxutils import escape

KINDS = {"name", "contact", "head", "line", "bullet", "para"}

def resume_blocks(r):
    s = lambda k: str(r.get(k) or "").strip()
    b = [("name", s("name")), ("contact", " | ".join(x for x in (s("email"), s("phone"), s("location"), s("linkedin")) if x))]
    if s("summary"): b += [("head", "PROFESSIONAL SUMMARY"), ("para", s("summary"))]
    if r.get("skills"): b += [("head", "SKILLS"), ("para", ", ".join(map(str, r["skills"])))]
    if r.get("experience"):
        b.append(("head", "EXPERIENCE"))
        for e in r["experience"]:
            b.append(("line", " - ".join(x for x in (str(e.get("title", "")), str(e.get("company", ""))) if x) + (f" | {e['dates']}" if e.get("dates") else "")))
            b += [("bullet", str(x)) for x in e.get("bullets", [])]
    if r.get("projects"):
        b.append(("head", "PROJECTS"))
        for p in r["projects"]:
            b.append(("line", str(p.get("name", ""))))
            b += [("bullet", str(x)) for x in p.get("bullets", [])]
    if r.get("education"):
        b.append(("head", "EDUCATION"))
        for e in r["education"]:
            b.append(("line", " - ".join(x for x in (str(e.get("degree", "")), str(e.get("school", ""))) if x) + (f" | {e['dates']}" if e.get("dates") else "")))
    if r.get("certifications"):
        b.append(("head", "CERTIFICATIONS")); b += [("bullet", str(x)) for x in r["certifications"]]
    return [(k, t) for k, t in b if t.strip()]

def text_blocks(text):
    out = []
    for line in text.splitlines():
        l = line.strip()
        if not l: continue
        if l.startswith(("- ", "* ")): out.append(("bullet", re.sub(r"[*_`]+", "", l[2:])))
        elif l.startswith("#"): out.append(("head", re.sub(r"[*_`]+", "", l.lstrip("# ")).upper()))
        else: out.append(("para", re.sub(r"[*_`]+", "", l)))
    return out

def clean_blocks(blocks):
    out = []
    for item in list(blocks)[:600]:
        if isinstance(item, (list, tuple)) and len(item) == 2 and item[0] in KINDS and str(item[1]).strip():
            out.append((item[0], str(item[1]).strip()[:600]))
    return out

def make_pdf(blocks):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle as P
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    st = {"name": P("n", fontName="Helvetica-Bold", fontSize=18, leading=22), "contact": P("c", fontName="Helvetica", fontSize=9.5, leading=13),
          "head": P("h", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=10, spaceAfter=3),
          "line": P("l", fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=4),
          "bullet": P("b", fontName="Helvetica", fontSize=10, leading=13, leftIndent=14, bulletIndent=4),
          "para": P("p", fontName="Helvetica", fontSize=10, leading=13)}
    fix = lambda t: escape(t.encode("latin-1", "replace").decode("latin-1"))
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=.75 * inch, rightMargin=.75 * inch, topMargin=.7 * inch, bottomMargin=.7 * inch)
    doc.build([Paragraph(fix(t), st[k], bulletText="-" if k == "bullet" else None) for k, t in blocks])
    return out.getvalue()

def make_docx(blocks):
    from docx import Document
    from docx.shared import Pt, Inches
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8); s.top_margin = s.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10.5)
    for k, t in blocks:
        if k == "bullet":
            doc.add_paragraph(t, style="List Bullet"); continue
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); run = p.add_run(t)
        if k == "name": run.bold = True; run.font.size = Pt(18)
        elif k == "head": run.bold = True; run.font.size = Pt(11); p.paragraph_format.space_before = Pt(10)
        elif k == "line": run.bold = True; p.paragraph_format.space_before = Pt(4)
        elif k == "contact": run.font.size = Pt(9.5)
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()
